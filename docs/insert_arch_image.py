# -*- coding: utf-8 -*-
"""把 architecture_v2.png 插入 02-系统设计.docx，替换文字架构图。"""
import os
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


def main():
    docs_dir = os.path.dirname(os.path.abspath(__file__))
    docx_path = os.path.join(docs_dir, "02-系统设计.docx")
    img_path = os.path.join(docs_dir, "architecture_v2.png")

    doc = Document(docx_path)

    # 找到标题 "1. 总体架构"、正文 "系统采用..."、文字架构图段落
    title_idx = None
    intro_idx = None
    arch_idx = None
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if "总体架构" in t:
            title_idx = i
        elif "三端分离" in t:
            intro_idx = i
        elif t.startswith("┌") or t.startswith("◆"):
            if arch_idx is None and intro_idx is not None:
                arch_idx = i

    if arch_idx is None and title_idx is not None:
        for i, p in enumerate(doc.paragraphs):
            if i > title_idx and len(p.text) > 100:
                arch_idx = i
                break

    print("title_idx:", title_idx, "intro_idx:", intro_idx, "arch_idx:", arch_idx)
    if arch_idx is None:
        print("未找到文字架构图段落")
        return

    # 删除文字架构图段落
    arch_p = doc.paragraphs[arch_idx]
    arch_p._element.getparent().remove(arch_p._element)

    # 在 intro 段落后插入图片
    insert_after = doc.paragraphs[intro_idx] if intro_idx is not None else doc.paragraphs[title_idx]
    p_el = insert_after._element.makeelement(qn('w:p'), {})
    insert_after._element.addnext(p_el)
    new_para = Paragraph(p_el, insert_after._parent)
    run = new_para.add_run()
    run.add_picture(img_path, width=Inches(6.5))

    doc.save(docx_path)
    print("OK: 新架构图已插入")


if __name__ == "__main__":
    main()
