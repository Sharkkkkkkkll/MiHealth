# -*- coding: utf-8 -*-
"""把 docx 里的架构图替换为新的 architecture_v2.png。"""
import os
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


def main():
    docs_dir = os.path.dirname(os.path.abspath(__file__))
    docx_path = os.path.join(docs_dir, "02-系统设计.docx")
    img_path = os.path.join(docs_dir, "architecture_v2.png")

    doc = Document(docx_path)

    # 找到包含图片的段落，删除图片 run
    removed = 0
    for p in doc.paragraphs:
        for el in list(p._element.findall(qn('w:r'))):
            if el.find(qn('w:drawing')) is not None:
                p._element.remove(el)
                removed += 1
    print("removed drawings:", removed)

    # 在 "系统采用「三端分离」架构" 段落后插入新图
    target = None
    for i, p in enumerate(doc.paragraphs):
        if "三端分离" in p.text:
            target = p
            break
    if target is None:
        print("未找到插入位置")
        return

    p_el = target._element.makeelement(qn('w:p'), {})
    target._element.addnext(p_el)
    new_para = Paragraph(p_el, target._parent)
    run = new_para.add_run()
    run.add_picture(img_path, width=Inches(6.5))

    tmp_path = os.path.join(docs_dir, "02-系统设计_new.docx")
    doc.save(tmp_path)
    print(f"OK: 已保存 {tmp_path}")
    # 尝试覆盖原文件（若 WPS 未锁定）
    try:
        os.replace(tmp_path, docx_path)
        print("OK: 已替换原文件")
    except Exception as e:
        print(f"原文件被占用，请手动替换：{e}")
        print(f"新文件在：{tmp_path}")


if __name__ == "__main__":
    main()
